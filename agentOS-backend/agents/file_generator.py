"""
Real File Output Agents - Generate actual files (PDF, PPTX, DOCX)
=================================================================

This makes AgentOS actually CREATE files, not just generate text.
Demo-worthy: "Make me a research paper" → Get actual PDF file.
"""
import os
import re
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

from core.llm import llm, rate_limit_delay
from config import settings


OUTPUT_DIR = Path("./outputs")
OUTPUT_DIR.mkdir(exist_ok=True)

# CSS styles for PDF/HTML output
HTML_CSS = """
body {
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    line-height: 1.7;
    color: #1a1a2e;
    max-width: 800px;
    margin: 0 auto;
    padding: 40px 20px;
    background: #fafafa;
}
h1 {font-size: 2.5em; font-weight: 700; color: #0f172a; margin-bottom: 0.5em; padding-bottom: 0.3em; border-bottom: 3px solid #3b82f6;}
h2 {font-size: 1.8em; font-weight: 600; color: #1e293b; margin-top: 1.5em; margin-bottom: 0.5em;}
h3 {font-size: 1.3em; font-weight: 600; color: #334155; margin-top: 1.2em;}
p {margin: 1em 0;}
code {font-family: 'Courier New', monospace; background: #f1f5f9; padding: 0.2em 0.4em; border-radius: 4px; font-size: 0.9em;}
pre {background: #1e293b; color: #e2e8f0; padding: 1em; border-radius: 8px; overflow-x: auto;}
pre code {background: none; padding: 0;}
blockquote {border-left: 4px solid #3b82f6; padding-left: 1em; margin: 1.5em 0; color: #64748b; font-style: italic;}
ul, ol {padding-left: 1.5em;}
li {margin: 0.5em 0;}
table {width: 100%; border-collapse: collapse; margin: 1.5em 0;}
th, td {border: 1px solid #e2e8f0; padding: 0.75em; text-align: left;}
th {background: #f1f5f9; font-weight: 600;}
img {max-width: 100%; height: auto; border-radius: 8px; margin: 1em 0;}
hr {border: none; border-top: 1px solid #e2e8f0; margin: 2em 0;}
a {color: #3b82f6; text-decoration: none;}
a:hover {text-decoration: underline;}
"""


class PDFGeneratorAgent:
    """Generates actual PDF files from markdown/HTML content."""
    
    def __init__(self):
        self._has_weasyprint = False
        try:
            import weasyprint
            # Test that it actually loads (requires GTK on Windows)
            from weasyprint import HTML
            HTML(string="<p>test</p>").write_pdf()
            self._has_weasyprint = True
        except Exception as e:
            # WeasyPrint requires GTK libs; fall back to HTML output
            print(f"[PDFGenerator] WeasyPrint not available: {e}. Using HTML fallback.")
            self._has_weasyprint = False
    
    async def run(self, content: str, task_id: str, title: str = "Document", format: str = "pdf") -> dict:
        print(f"[PDFGenerator] Creating {format} for task: {task_id}")
        
        if self._has_weasyprint:
            return await self._generate_weasyprint(content, task_id, title, format)
        else:
            return await self._generate_html(content, task_id, title)
    
    async def _generate_weasyprint(self, content: str, task_id: str, title: str, format: str) -> dict:
        from weasyprint import HTML
        
        html_content = self._create_html_document(content, title)
        pdf_bytes = HTML(string=html_content).write_pdf()
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{task_id}_{timestamp}.{format}"
        filepath = OUTPUT_DIR / filename
        
        with open(filepath, "wb") as f:
            f.write(pdf_bytes)
        
        print(f"[PDFGenerator] Saved to: {filepath}")
        
        return {
            "success": True,
            "file_path": str(filepath),
            "filename": filename,
            "file_size": len(pdf_bytes),
            "download_url": f"/api/files/{filename}",
            "format": format,
            "title": title,
            "created_at": datetime.now().isoformat(),
        }
    
    async def _generate_html(self, content: str, task_id: str, title: str) -> dict:
        html_content = self._create_html_document(content, title)
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{task_id}_{timestamp}.html"
        filepath = OUTPUT_DIR / filename
        
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(html_content)
        
        return {
            "success": True,
            "file_path": str(filepath),
            "filename": filename,
            "file_size": len(html_content.encode()),
            "download_url": f"/api/files/{filename}",
            "format": "html",
            "title": title,
            "created_at": datetime.now().isoformat(),
            "note": "HTML generated. Install weasyprint for PDF.",
        }
    
    def _create_html_document(self, md: str, title: str) -> str:
        import markdown
        html_body = markdown.markdown(md, extensions=['extra', 'codehilite', 'tables', 'fenced_code'])
        
        return f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <style>{HTML_CSS}</style>
</head>
<body>
    <div class="document-header">
        <h1>{title}</h1>
    </div>
    {html_body}
    <div class="document-footer">
        <p>Generated by AgentOS on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
    </div>
</body>
</html>"""


class PPTXGeneratorAgent:
    """Generates actual PowerPoint (.pptx) files."""
    
    def __init__(self):
        self._has_pptx = False
        try:
            from pptx import Presentation
            self._has_pptx = True
        except ImportError:
            pass
    
    async def run(self, content: str, task_id: str, title: str = "Presentation") -> dict:
        print(f"[PPTXGenerator] Creating presentation for task: {task_id}")
        
        if not self._has_pptx:
            return {"success": False, "error": "python-pptx not installed"}
        
        return await self._generate_pptx(content, task_id, title)
    
    async def _generate_pptx(self, content: str, task_id: str, title: str) -> dict:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        # Title slide
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = title
        title_slide.placeholders[1].text = f"Generated by AgentOS\n{datetime.now().strftime('%B %Y')}"
        
        # Parse and add content slides
        slides = self._parse_slides(content)
        for slide_data in slides:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = slide_data.get("title", "Slide")
            
            tf = slide.placeholders[1].text_frame
            tf.clear()
            for bullet in slide_data.get("bullets", []):
                p = tf.add_paragraph()
                p.text = bullet
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{task_id}_{timestamp}.pptx"
        filepath = OUTPUT_DIR / filename
        
        prs.save(str(filepath))
        
        return {
            "success": True,
            "file_path": str(filepath),
            "filename": filename,
            "file_size": filepath.stat().st_size,
            "download_url": f"/api/files/{filename}",
            "format": "pptx",
            "title": title,
            "slide_count": len(slides) + 1,
            "created_at": datetime.now().isoformat(),
        }
    
    def _parse_slides(self, content: str) -> list:
        slides = []
        parts = re.split(r'^---\s*$|^##\s+', content, flags=re.MULTILINE)
        
        for part in parts:
            part = part.strip()
            if not part:
                continue
            
            lines = part.split('\n')
            title = lines[0].strip() if lines else "Slide"
            bullets = []
            
            for line in lines[1:]:
                line = line.strip()
                if line.startswith('- ') or line.startswith('* '):
                    bullets.append(line[2:])
                elif line:
                    bullets.append(line)
            
            if title or bullets:
                slides.append({"title": title, "bullets": bullets})
        
        return slides


class DOCXGeneratorAgent:
    """Generates actual Word (.docx) files."""
    
    def __init__(self):
        self._has_docx = False
        try:
            from docx import Document
            self._has_docx = True
        except ImportError:
            pass
    
    async def run(self, content: str, task_id: str, title: str = "Document") -> dict:
        print(f"[DOCXGenerator] Creating document for task: {task_id}")
        
        if not self._has_docx:
            return {"success": False, "error": "python-docx not installed"}
        
        return await self._generate_docx(content, task_id, title)
    
    async def _generate_docx(self, content: str, task_id: str, title: str) -> dict:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        doc = Document()
        
        # Title
        title_para = doc.add_heading(title, 0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()
        
        # Process content
        lines = content.split('\n')
        for line in lines:
            line = line.rstrip()
            
            if line.startswith('# '):
                doc.add_heading(line[2:], 0)
            elif line.startswith('## '):
                doc.add_heading(line[3:], 1)
            elif line.startswith('### '):
                doc.add_heading(line[4:], 2)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.strip():
                doc.add_paragraph(line)
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{task_id}_{timestamp}.docx"
        filepath = OUTPUT_DIR / filename
        
        doc.save(str(filepath))
        
        return {
            "success": True,
            "file_path": str(filepath),
            "filename": filename,
            "file_size": filepath.stat().st_size,
            "download_url": f"/api/files/{filename}",
            "format": "docx",
            "title": title,
            "created_at": datetime.now().isoformat(),
        }


def get_file_generator(output_format: str):
    generators = {"pdf": PDFGeneratorAgent, "pptx": PPTXGeneratorAgent, "docx": DOCXGeneratorAgent, "html": PDFGeneratorAgent}
    generator_class = generators.get(output_format.lower())
    if generator_class is None:
        raise ValueError(f"Unknown format: {output_format}")
    return generator_class()


# Singleton instances
pdf_generator = PDFGeneratorAgent()
pptx_generator = PPTXGeneratorAgent()
docx_generator = DOCXGeneratorAgent()
